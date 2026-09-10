# -*- coding: utf-8 -*-
"""xlsx_deliverable.py —— Excel 交付物「行列位」体检与安全填值

背景（真实翻车记录）
-------------------
数学建模赛题常给 resultN.xlsx **模板**：第 1 行是时间表头、第 1 列是把手/对象标签，
数据区从「第 2 行 × 第 2 列」开始。把数据从第 1 行或第 1 列开始写，会
① 覆盖表头/标签行、② 整块错位一格、③ 末行/末列留空 —— 数值全对但**结构全错**，
且按标签读数会把对象整体错位一个（例如"第51节"读成"第50节"）。
本项目曾同时在 result1/result2/result4 三个文件上犯此错，靠独立复核才发现。

用法
----
1) 看一眼模板的「行列契约」（写数据之前必须做）：
     python xlsx_deliverable.py probe 附件/result4.xlsx
2) 填完数据后体检（推荐把契约显式写进命令，便于复现）：
     python xlsx_deliverable.py check out/result4.xlsx \
         --sheet 位置 --header-row 1 --label-col 1 \
         --data-row-first 2 --data-row-last 449 \
         --data-col-first 2 --data-col-last 202
3) 只想自动体检（从文件本身推断契约）：
     python xlsx_deliverable.py check out/result4.xlsx --sheet 位置 --auto

退出码：0 = 全部 PASS；1 = 有 FAIL；2 = 用法/依赖错误。

依赖：openpyxl（缺失则提示安装；本工具不需要 pandas）。
"""
import argparse
import sys

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print("[xlsx_deliverable] 需要 openpyxl：pip install openpyxl", file=sys.stderr)
    sys.exit(2)

RESULT = []


def _rec(ok, tag, sheet, item, expect, got):
    RESULT.append((bool(ok), sheet, item, expect, got))
    flag = "PASS" if ok else "FAIL"
    print(f"  [{flag}] {sheet:8s} | {item:38s} | 期望: {expect:24s} | 实测: {got}")


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def probe_sheet(ws, max_scan_row=3000, max_scan_col=1000):
    """自动推断一个工作表的「行列契约」"""
    nrow = min(ws.max_row, max_scan_row)
    ncol = min(ws.max_column, max_scan_col)

    def cell(r, c):
        return ws.cell(row=r, column=c).value

    # 1) 表头行：第 1 列空(或非标签) 且右侧至少 2 个非空
    header_row = None
    for r in range(1, min(nrow, 20) + 1):
        right = sum(0 if _blank(cell(r, c)) else 1 for c in range(2, min(ncol, 40) + 1))
        if right >= 2 and _blank(cell(r, 1)):
            header_row = r
            break

    # 2) 标签列：表头行下方、第 1 列连续非空的列（一般为 1）
    label_col = 1

    # 3) 数据区：表头行之后 + 标签列右侧
    first_data_row = (header_row + 1) if header_row else 1
    last_data_row = nrow
    while last_data_row > first_data_row:
        if any(not _blank(cell(last_data_row, c)) for c in range(label_col + 1, ncol + 1)):
            break
        last_data_row -= 1
    last_data_col = ncol
    while last_data_col > label_col + 1:
        if any(not _blank(cell(r, last_data_col)) for r in range(first_data_row, last_data_row + 1)):
            break
        last_data_col -= 1

    # 空模板兜底：模板通常只有表头+标签、数据区全空。此时上面的"收缩"会把数据区塌缩成 1×1，
    # 必须回退到工作表的**结构范围**（ws.max_row/max_column 反映模板已设定的行列），
    # 否则给出的"填值约定"会是错的（这正是本工具存在的首要原因）。
    is_empty = not any(not _blank(cell(r, c))
                       for r in range(first_data_row, last_data_row + 1)
                       for c in range(label_col + 1, last_data_col + 1))
    if is_empty:
        last_data_row = nrow
        last_data_col = ncol

    return {
        "sheet": ws.title,
        "dims": (ws.max_row, ws.max_column),
        "is_empty": is_empty,
        "header_row": header_row,
        "label_col": label_col,
        "data_row_first": first_data_row,
        "data_row_last": last_data_row,
        "data_col_first": label_col + 1,
        "data_col_last": last_data_col,
        "header_samples": [cell(header_row, c) for c in (2, 3, last_data_col)] if header_row else [],
        "label_samples": [cell(r, 1) for r in (first_data_row, first_data_row + 1, last_data_row)],
    }


def cmd_probe(args):
    wb = load_workbook(args.file)
    sheets = [args.sheet] if args.sheet else wb.sheetnames
    for name in sheets:
        ws = wb[name]
        c = probe_sheet(ws)
        print("=" * 96)
        print(f"工作表 [{c['sheet']}]  尺寸 {c['dims'][0]} 行 × {c['dims'][1]} 列"
              + ("   [空模板: 数据区无值, 下表按结构范围推断]" if c.get("is_empty") else ""))
        print(f"  表头行      : 第 {c['header_row']} 行   样例(第2/3/末列) = {c['header_samples']}")
        print(f"  标签列      : 第 {c['label_col']} 列   样例(首/次/末行) = {c['label_samples']}")
        print(f"  数据行区间  : 第 {c['data_row_first']}..{c['data_row_last']} 行"
              f"  (行数 {c['data_row_last'] - c['data_row_first'] + 1})")
        print(f"  数据列区间  : 第 {c['data_col_first']}..{c['data_col_last']} 列"
              f"  (列数 {c['data_col_last'] - c['data_col_first'] + 1})")
        print("\n  >>> 填值约定（照抄到你的写入代码）:")
        print(f"      第 {c['data_row_first']} 行 = 第 1 个数据（表头行第 {c['header_row']} 行不得覆盖）")
        print(f"      第 {c['data_col_first']} 列 = 第 1 个时间/维度（标签列第 {c['label_col']} 列不得覆盖）")
        print(f"      对象 i（0-based）→ 行 = {c['data_row_first']} + i * K（K=每对象占行数，自行填）")
    return 0


def cmd_check(args):
    wb = load_workbook(args.file)
    sheets = [args.sheet] if args.sheet else wb.sheetnames
    for name in sheets:
        ws = wb[name]
        auto = probe_sheet(ws)
        header_row = args.header_row if args.header_row is not None else auto["header_row"]
        label_col = args.label_col if args.label_col is not None else auto["label_col"]
        r1 = args.data_row_first if args.data_row_first is not None else auto["data_row_first"]
        r2 = args.data_row_last if args.data_row_last is not None else auto["data_row_last"]
        c1 = args.data_col_first if args.data_col_first is not None else auto["data_col_first"]
        c2 = args.data_col_last if args.data_col_last is not None else auto["data_col_last"]

        print("=" * 96)
        print(f"工作表 [{name}]  契约: 表头行={header_row}, 标签列={label_col}, "
              f"数据 {r1}..{r2} 行 × {c1}..{c2} 列")

        # ① 表头行未被覆盖：表头行第 2 列起应有非空且不是纯数字
        if header_row:
            hdr_vals = [ws.cell(row=header_row, column=c).value for c in range(c1, min(c2, c1 + 3) + 1)]
            ok = any(not _blank(v) for v in hdr_vals) and \
                 any(isinstance(v, str) and not _is_numberish(v) for v in hdr_vals)
            _rec(ok, "header-intact", name, f"表头行第 {header_row} 行未被覆盖",
                 "含时间/维度文字", f"{hdr_vals[:3]}")

        # ② 标签列未被覆盖：数据首行/末行的第 label_col 列应为文字标签
        lab_first = ws.cell(row=r1, column=label_col).value
        lab_last = ws.cell(row=r2, column=label_col).value
        ok = isinstance(lab_first, str) and not _is_numberish(lab_first) and \
             isinstance(lab_last, str) and not _is_numberish(lab_last)
        _rec(ok, "labels-intact", name, f"标签列第 {label_col} 列首/末行是文字标签",
             "文字（如 龙头 / 龙尾（后））", f"{lab_first!r} / {lab_last!r}")

        # ③ 数据区不与表头/标签列重叠（结构性红线）
        ok = (r1 > (header_row or 0)) and (c1 > label_col)
        _rec(ok, "no-overlap", name, "数据区起始行/列在表头与标签之后",
             f"行>{header_row}, 列>{label_col}", f"行{r1}, 列{c1}")

        # ④ 数据区无空洞
        holes = 0
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if _blank(ws.cell(row=r, column=c).value):
                    holes += 1
        _rec(holes == 0, "no-holes", name, f"数据区 {r1}..{r2} × {c1}..{c2} 无空值",
             "0 个空值", f"{holes} 个空值")

        # ⑤ 末行/末列非空（防"整体上移/左移一格"）
        last_cell = ws.cell(row=r2, column=c2).value
        _rec(not _blank(last_cell), "corner-filled", name, f"数据区右下角 ({r2},{c2}) 非空",
             "非空", f"{last_cell!r}")
        if r2 < ws.max_row:
            extra = [ws.cell(row=r, column=c).value
                     for r in range(r2 + 1, min(ws.max_row, r2 + 3) + 1)
                     for c in range(c1, c2 + 1)]
            n_extra = sum(0 if _blank(v) else 1 for v in extra)
            _rec(n_extra == 0, "no-spill-below", name,
                 f"第 {r2} 行以下无多余数据", "空白", f"{n_extra} 个非空")
        if c2 < ws.max_column:
            extra = [ws.cell(row=r, column=c).value
                     for c in range(c2 + 1, min(ws.max_column, c2 + 3) + 1)
                     for r in range(r1, r2 + 1)]
            n_extra = sum(0 if _blank(v) else 1 for v in extra)
            _rec(n_extra == 0, "no-spill-right", name,
                 f"第 {c2} 列右侧无多余数据", "空白", f"{n_extra} 个非空")

    n_fail = sum(1 for ok, *_ in RESULT if not ok)
    print("\n" + "=" * 96)
    print(f"汇总: 共 {len(RESULT)} 项, PASS {len(RESULT) - n_fail}, FAIL {n_fail}")
    if n_fail:
        print("FAIL 明细:")
        for ok, sh, item, exp, got in RESULT:
            if not ok:
                print(f"  × [{sh}] {item}: 期望 {exp}, 实测 {got}")
        print("总判定: FAIL —— 数值可能全对，但交付物结构不合格，评委按标签读数会错位。")
        return 1
    print("总判定: PASS —— 行列位合格，标签与数据一一对齐。")
    return 0


def _is_numberish(v):
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, str):
        s = v.strip().replace("-", "").replace(".", "").replace("e", "").replace("E", "")
        return s.isdigit()
    return False


def main():
    ap = argparse.ArgumentParser(description="Excel 交付物行列位体检")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p1 = sub.add_parser("probe", help="推断模板的行列契约（写数据前先跑）")
    p1.add_argument("file")
    p1.add_argument("--sheet", default=None)
    p1.set_defaults(func=cmd_probe)

    p2 = sub.add_parser("check", help="体检已填好的交付物")
    p2.add_argument("file")
    p2.add_argument("--sheet", default=None)
    p2.add_argument("--header-row", type=int, default=None)
    p2.add_argument("--label-col", type=int, default=None)
    p2.add_argument("--data-row-first", type=int, default=None)
    p2.add_argument("--data-row-last", type=int, default=None)
    p2.add_argument("--data-col-first", type=int, default=None)
    p2.add_argument("--data-col-last", type=int, default=None)
    p2.add_argument("--auto", action="store_true", help="全部参数自动推断")
    p2.set_defaults(func=cmd_check)

    args = ap.parse_args()
    sys.exit(args.func(args))


if __name__ == "__main__":
    main()
